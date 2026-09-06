from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION

OUT = '/Users/balu/Documents/Codex/2026-08-14/i-n/outputs/Food-Delivery-App-Proposal-and-Budget.docx'

NAVY = '17324D'
BLUE = '2E74B5'
MUTED = '596775'
PALE = 'EAF1F7'
CORAL = 'FF5A4F'
INK = '1F2933'

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tc_pr.append(mar)
    for side, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW'); tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), '9360'); tbl_w.set(qn('w:type'), 'dxa')
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa'); tbl_pr.append(ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tcw = tc_pr.first_child_found_in('w:tcW')
            if tcw is None:
                tcw = OxmlElement('w:tcW'); tc_pr.append(tcw)
            tcw.set(qn('w:w'), str(width)); tcw.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_rule(p, color=NAVY):
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12'); bottom.set(qn('w:space'), '8'); bottom.set(qn('w:color'), color)
    borders.append(bottom); p_pr.append(borders)

def add_text(p, text, size=11, color=INK, bold=False, italic=False):
    r = p.add_run(text); r.font.name = 'Calibri'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color); r.bold = bold; r.italic = italic
    return r

def para(doc, text='', size=11, color=INK, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.1
    if align is not None: p.alignment = align
    if text: add_text(p, text, size, color, bold)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.12
    add_text(p, text, 10.5)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(16 if level == 1 else 10); p.paragraph_format.space_after = Pt(7)
    add_text(p, text, 16 if level == 1 else 12.5, BLUE if level == 1 else NAVY, True)
    return p

def set_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, 'CUSTOM FOOD DELIVERY PLATFORM  |  COMMERCIAL PROPOSAL', 8.5, MUTED, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, 'Confidential commercial proposal  •  Amounts in INR  •  Page ', 8.5, MUTED)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.85)
section.header_distance = Inches(.35); section.footer_distance = Inches(.35)
set_header_footer(section)

styles = doc.styles
styles['Normal'].font.name = 'Calibri'; styles['Normal'].font.size = Pt(11)
for name in ['Title', 'Subtitle']:
    styles[name].font.name = 'Calibri'

# Cover / proposal masthead
p = para(doc, 'COMMERCIAL PROPOSAL', 10, CORAL, True, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
p = para(doc, 'Custom Food Delivery\nMobile Application', 27, NAVY, True, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.line_spacing = 0.95
para(doc, 'Android + iOS customer apps, delivery-partner workflow, restaurant operations and admin controls', 12.5, MUTED, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)

meta = doc.add_table(rows=2, cols=2)
set_table_width(meta, [4680, 4680])
for row in meta.rows:
    for cell in row.cells: set_cell_shading(cell, PALE)
labels = [('Prepared for', 'Client / Business Name'), ('Prepared by', 'Your Company / Name'), ('Proposal date', '14 August 2026'), ('Validity', '30 days from proposal date')]
for cell, (label, value) in zip([c for r in meta.rows for c in r.cells], labels):
    p = cell.paragraphs[0]; add_text(p, label.upper() + '\n', 8.5, MUTED, True); add_text(p, value, 10.5, NAVY, True)

para(doc, '', after=10)
callout = doc.add_table(rows=1, cols=1); set_table_width(callout, [9360]); set_cell_shading(callout.cell(0,0), NAVY)
p = callout.cell(0,0).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, 'YEAR 1 PROJECT INVESTMENT: ₹10,00,000 + applicable GST', 16, 'FFFFFF', True)
p = callout.cell(0,0).add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, 'Full source-code and business-asset handover after final payment', 10, 'D9E8F5')

heading(doc, 'Executive summary')
para(doc, 'This proposal covers the design, development, testing and launch support for a premium, custom food-delivery platform. The finished solution will be branded exclusively for the client and handed over with the source code, database access and administrative access after final payment.')
para(doc, 'This is not a rented annual marketplace service. The client will own the software and business assets, while optional future support can be requested whenever needed.', after=10)

heading(doc, 'Included solution')
for item in [
    'Customer mobile application for Android and iPhone: discovery, menus, cart, checkout, offers, notifications, order history, ratings and delivery-status tracking.',
    'Delivery-partner application: availability, order acceptance, pickup/delivery status, navigation handoff, earnings and delivery history.',
    'Restaurant dashboard: restaurant profile, menus, item availability, incoming orders, preparation status, order history and sales information.',
    'Admin dashboard: users, restaurants, delivery partners, orders, delivery zones, commissions, banners, coupons, reports and operational controls.',
    'Secure backend: database, login, APIs, role-based access, configuration, backups and integration-ready architecture.',
    'Custom UI/UX, logo, colour system, app icon, launch testing and Google Play / Apple App Store submission assistance.'
]: bullet(doc, item)

heading(doc, 'Development budget')
budget = doc.add_table(rows=1, cols=2); set_table_width(budget, [6900, 2460])
for cell, text in zip(budget.rows[0].cells, ['Deliverable', 'Amount (₹)']):
    set_cell_shading(cell, PALE); p = cell.paragraphs[0]; add_text(p, text, 10, NAVY, True); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if text.startswith('Amount') else WD_ALIGN_PARAGRAPH.LEFT
items = [
    ('UI/UX, branding and application logo', '80,000'),
    ('Customer application — Android and iOS', '2,80,000'),
    ('Delivery-partner application', '1,50,000'),
    ('Restaurant partner and super-admin dashboards', '1,10,000'),
    ('Backend, database, APIs and authentication', '2,00,000'),
    ('Payments, Maps integration point, tracking and notifications', '1,00,000'),
    ('Testing, release support and complete source-code handover', '80,000'),
]
for label, amount in items:
    cells = budget.add_row().cells; add_text(cells[0].paragraphs[0], label, 10.3); p = cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(p, amount, 10.3, INK, True)
cells = budget.add_row().cells
for c in cells: set_cell_shading(c, NAVY)
add_text(cells[0].paragraphs[0], 'TOTAL ONE-TIME DEVELOPMENT COST', 10.5, 'FFFFFF', True)
p = cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(p, '10,00,000', 10.5, 'FFFFFF', True)

heading(doc, 'Post-launch support and third-party costs')
para(doc, 'The project includes 60 days of post-launch bug-fix support. After this period, support and minor changes are charged only when requested.', after=5)
support = doc.add_table(rows=1, cols=2); set_table_width(support, [6900, 2460])
for cell, text in zip(support.rows[0].cells, ['Support item', 'Commercial terms']):
    set_cell_shading(cell, PALE); add_text(cell.paragraphs[0], text, 10, NAVY, True)
for label, value in [
    ('Post-launch bug-fix support', 'Included for 60 days'),
    ('Support request / minor change after free period', '₹15,000 per request'),
    ('Major new feature or integration', 'Quoted separately after approval'),
    ('Payment gateway, SMS/OTP, Maps, hosting and cloud usage', 'Client pays actual usage charges'),
    ('Apple Developer and Google Play accounts', 'Client-owned accounts; fees paid directly by client'),
]:
    cells = support.add_row().cells; add_text(cells[0].paragraphs[0], label, 10.3); p = cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_text(p, value, 10.1, INK, True)

doc.add_page_break()
heading(doc, 'Delivery plan and payment schedule')
para(doc, 'Estimated delivery period: 20–24 weeks from confirmation, initial payment, and approval of the branding and required business information.')
schedule = doc.add_table(rows=1, cols=3); set_table_width(schedule, [1650, 5260, 2450])
for c, t in zip(schedule.rows[0].cells, ['Milestone', 'Outcome', 'Amount (₹)']):
    set_cell_shading(c, PALE); add_text(c.paragraphs[0], t, 10, NAVY, True)
for milestone, outcome, amount in [
    ('1', 'Project confirmation, planning and product discovery', '2,00,000'),
    ('2', 'Branding, UI/UX approval and customer-app development', '2,50,000'),
    ('3', 'Rider app, dashboards, backend and integrations', '3,00,000'),
    ('4', 'Testing, deployment and store-submission preparation', '1,50,000'),
    ('5', 'Final launch and complete handover', '1,00,000'),
]:
    cells = schedule.add_row().cells
    for i, value in enumerate([milestone, outcome, amount]):
        p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT; add_text(p, value, 10.1, INK, i in [0,2])

heading(doc, 'Assumptions and commercial terms')
for item in [
    'The initial launch covers one primary city and standard food-delivery operations. Expansion to additional cities can be configured after launch.',
    'The client will provide business information, restaurant/menu content, legal documents, payment-gateway onboarding details and store-account credentials in a timely manner.',
    'Client-owned accounts will be used for payments, cloud hosting, Apple, Google, SMS/OTP and Maps services. This preserves ownership and avoids dependency on a third party.',
    'The client receives the custom source code, branding assets, database access and admin credentials after final development payment is received.',
    'Support covers issue diagnosis and agreed minor changes. Materially new features, redesigns, additional integrations or scope changes require a separate written estimate.',
    'Amounts exclude applicable GST and third-party usage charges.'
]: bullet(doc, item)

heading(doc, 'Acceptance')
para(doc, 'Acceptance of this proposal confirms the scope, payment schedule and commercial terms stated above. Development begins after the first milestone payment and written approval to proceed.')
para(doc, '', after=20)
sig = doc.add_table(rows=2, cols=2); set_table_width(sig, [4680, 4680])
for cell, text in zip(sig.rows[0].cells, ['CLIENT AUTHORISED SIGNATORY', 'SERVICE PROVIDER']):
    set_cell_shading(cell, PALE); p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_text(p, text, 9, NAVY, True)
for cell in sig.rows[1].cells:
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(20); add_text(p, '\nName: __________________________\nSignature / Date: __________________', 10, MUTED)

doc.save(OUT)
print(OUT)
